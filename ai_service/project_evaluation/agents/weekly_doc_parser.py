import json
import os
from pathlib import Path

from docx import Document
from langchain_core.prompts import PromptTemplate
from langchain_groq import ChatGroq
from pypdf import PdfReader


def extract_weekly_report_text(report_path):
    path = Path(report_path)

    if not path.is_file():
        raise FileNotFoundError(f"Weekly report not found: {path}")

    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8-sig")
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if suffix == ".docx":
        document = Document(str(path))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join(paragraphs + table_rows)

    raise ValueError("Weekly report must be a PDF, DOCX, or TXT file")


def parse_weekly_report(report_path):
    raw_text = extract_weekly_report_text(report_path)
    return parse_weekly_report_text(raw_text)


def parse_weekly_report_text(raw_text):
    if not isinstance(raw_text, str):
        raise TypeError("Weekly report text must be a string")
    if not raw_text.strip():
        raise ValueError("Weekly report contains no extractable text")

    if not os.getenv("GROQ_API_KEY"):
        raise RuntimeError("GROQ_API_KEY is not set")

    prompt = PromptTemplate.from_template("""
You extract information from a student's weekly progress report.

Extract only claims the student explicitly makes. Do not assess whether the
claims are true, complete, or supported by code. Do not classify claims as
SUPPORTED, PARTIAL, or UNSUPPORTED, and do not infer misconduct.
Rewrite claims as concise action or issue statements while preserving their
meaning. Put work already done in workCompleted, difficulties in problemsFaced,
and future intentions in plannedWork. Summarize the student's reported state
in currentStatus using only information in the report. Use an empty string if
the report does not describe a current status.

REPORT:
{raw_text}

Return ONLY valid JSON with exactly this structure:
{{
    "workCompleted": ["..."],
    "problemsFaced": ["..."],
    "plannedWork": ["..."],
    "currentStatus": "..."
}}
""")

    llm = ChatGroq(
        model=os.getenv("GROQ_MODEL", "openai/gpt-oss-20b"),
        temperature=0,
    )
    response = llm.invoke(prompt.format(raw_text=raw_text))

    try:
        parsed = json.loads(response.content)
    except json.JSONDecodeError as exc:
        raise ValueError("Groq returned invalid JSON for the weekly report") from exc

    result = {
        "workCompleted": parsed.get("workCompleted", []),
        "problemsFaced": parsed.get("problemsFaced", []),
        "plannedWork": parsed.get("plannedWork", []),
        "currentStatus": parsed.get("currentStatus", ""),
        "rawText": raw_text,
    }
    for field in ("workCompleted", "problemsFaced", "plannedWork"):
        if not isinstance(result[field], list) or not all(
            isinstance(item, str) for item in result[field]
        ):
            raise ValueError(f"Groq returned an invalid {field} value")
    if not isinstance(result["currentStatus"], str):
        raise ValueError("Groq returned an invalid currentStatus value")

    return result